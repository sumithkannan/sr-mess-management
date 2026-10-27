from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from passlib.context import CryptContext
from openpyxl import Workbook, load_workbook
from io import BytesIO
from ..dependencies import get_db, get_current_user, admin_required
from ..models.user import User
from ..schemas.user import UserCreate, UserUpdate, UserResponse

router = APIRouter(prefix="/api/users", tags=["Users"])
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

@router.get("/active")
def list_active_users(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    users = db.query(User).filter(User.is_active == True).all()
    return [{"id": u.id, "name": u.name, "username": u.username} for u in users]

@router.get("/bulk/template")
def download_template(admin=Depends(admin_required)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    ws.append(["Username", "Name", "Email"])
    ws.append(["example_user", "Example Name", "example@email.com"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=users_bulk_template.xlsx"}
    )

@router.post("/bulk")
async def bulk_add_users(file: UploadFile = File(...), db: Session = Depends(get_db), admin=Depends(admin_required)):
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Only .xlsx files are allowed")
    
    try:
        content = await file.read()
        wb = load_workbook(BytesIO(content))
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid Excel file")
    
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) == 0:
        raise HTTPException(status_code=400, detail="Empty file")
    
    headers = [str(h).strip().lower() if h else '' for h in rows[0]]
    if headers != ['username', 'name', 'email']:
        raise HTTPException(status_code=400, detail="Template mismatch. Expected columns: Username, Name, Email")
    
    data_rows = rows[1:]
    if len(data_rows) == 0:
        raise HTTPException(status_code=400, detail="No data rows found")
    
    existing_usernames = set(u.username.lower() for u in db.query(User.username).all())
    existing_emails = set(u.email.lower() for u in db.query(User.email).filter(User.email != None).all())
    
    seen_usernames = set()
    seen_emails = set()
    
    for idx, row in enumerate(data_rows, start=2):
        username = str(row[0]).strip() if row[0] else ''
        name = str(row[1]).strip() if row[1] else ''
        email = str(row[2]).strip() if row[2] else ''
        email = email if email and email.lower() != 'none' else None
        
        if not username or not name:
            raise HTTPException(status_code=400, detail=f"Row {idx}: Username and Name are required")
        if username.lower() in existing_usernames or username.lower() in seen_usernames:
            raise HTTPException(status_code=400, detail=f"Row {idx}: Duplicate username '{username}'")
        if email and (email.lower() in existing_emails or email.lower() in seen_emails):
            raise HTTPException(status_code=400, detail=f"Row {idx}: Duplicate email '{email}'")
        
        seen_usernames.add(username.lower())
        if email:
            seen_emails.add(email.lower())
    
    added = 0
    for row in data_rows:
        username = str(row[0]).strip()
        name = str(row[1]).strip()
        email = str(row[2]).strip() if row[2] else ''
        email = email if email and email.lower() != 'none' else None
        user = User(
            username=username,
            name=name,
            email=email,
            password_hash=pwd_context.hash("PasswordToBeChanged"),
            role="user"
        )
        db.add(user)
        added += 1
    
    db.commit()
    return {"added": added, "message": f"Successfully added {added} user(s)"}

@router.get("", response_model=list[UserResponse])
def list_users(db: Session = Depends(get_db), admin=Depends(admin_required)):
    return db.query(User).all()

@router.post("", response_model=UserResponse)
def create_user(req: UserCreate, db: Session = Depends(get_db), admin=Depends(admin_required)):
    if db.query(User).filter(User.username == req.username).first():
        raise HTTPException(status_code=400, detail="Username already exists")
    password = req.password if req.password else "PasswordToBeChanged"
    email = req.email.strip() if req.email and req.email.strip() else None
    phone = req.phone.strip() if req.phone and req.phone.strip() else None
    user = User(
        username=req.username, name=req.name, email=email, phone=phone,
        password_hash=pwd_context.hash(password)
    )
    db.add(user); db.commit(); db.refresh(user)
    return UserResponse.model_validate(user)

@router.put("/{user_id}", response_model=UserResponse)
def update_user(user_id: int, req: UserUpdate, db: Session = Depends(get_db), admin=Depends(admin_required)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    if req.name is not None: user.name = req.name
    if req.email is not None: user.email = req.email.strip() if req.email.strip() else None
    if req.phone is not None: user.phone = req.phone.strip() if req.phone.strip() else None
    if req.password: user.password_hash = pwd_context.hash(req.password)
    db.commit(); db.refresh(user)
    return UserResponse.model_validate(user)

@router.patch("/{user_id}/suspend")
def suspend_user(user_id: int, db: Session = Depends(get_db), admin=Depends(admin_required)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(status_code=404)
    user.is_active = False; db.commit()
    return {"message": "User suspended"}

@router.patch("/{user_id}/activate")
def activate_user(user_id: int, db: Session = Depends(get_db), admin=Depends(admin_required)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(status_code=404)
    user.is_active = True; db.commit()
    return {"message": "User activated"}

@router.patch("/{user_id}/role")
def change_role(user_id: int, role: str, db: Session = Depends(get_db), admin=Depends(admin_required)):
    if role not in ["admin", "user"]:
        raise HTTPException(status_code=400, detail="Invalid role")
    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(status_code=404)
    user.role = role; db.commit()
    return {"message": f"Role changed to {role}"}

@router.delete("/{user_id}")
def delete_user(user_id: int, db: Session = Depends(get_db), admin=Depends(admin_required)):
    user = db.query(User).filter(User.id == user_id).first()
    if not user: raise HTTPException(status_code=404)
    db.delete(user); db.commit()
    return {"message": "User deleted"}
